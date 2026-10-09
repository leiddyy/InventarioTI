import io
import os
from datetime import datetime
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from reportlab.lib.pagesizes import letter, landscape
from reportlab.lib import colors
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle

def gerar_excel_inventario(equipamentos) -> io.BytesIO:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Inventário de TI DIA"

    # Styles
    navy_fill = PatternFill(start_color="162052", end_color="162052", fill_type="solid")
    gold_fill = PatternFill(start_color="FDE68A", end_color="FDE68A", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    title_font = Font(name="Calibri", size=16, bold=True, color="162052")
    subtitle_font = Font(name="Calibri", size=11, bold=True, color="D97706")
    bold_font = Font(name="Calibri", size=10, bold=True)
    normal_font = Font(name="Calibri", size=10)
    
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    # Title rows
    ws.append(["DIA DISTRIBUIÇÃO E IMPORTAÇÃO AFOGADOS LTDA"])
    ws.cell(row=1, column=1).font = title_font
    ws.append(["RELATÓRIO OFICIAL DE INVENTÁRIO DE TI"])
    ws.cell(row=2, column=1).font = subtitle_font
    ws.append([f"Data de Emissão: {datetime.now().strftime('%d/%m/%Y às %H:%M')} | Total de Ativos: {len(equipamentos)}"])
    ws.cell(row=3, column=1).font = Font(name="Calibri", size=10, italic=True, color="64748B")
    ws.append([]) # Empty row

    # Header columns
    headers = [
        "Patrimônio", "Categoria", "Tipo", "Marca / Modelo",
        "Nº de Série", "Setor em Uso", "Responsável Atual",
        "Qtd", "Condição", "Status", "Valor (R$)", "Especificações / Hardware"
    ]
    ws.append(headers)
    
    header_row_idx = 5
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=header_row_idx, column=col_num)
        cell.fill = navy_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    # Data rows
    total_valor = 0.0
    for eq in equipamentos:
        valor_val = eq.valor or 0.0
        total_valor += valor_val * eq.quantidade

        specs = []
        if eq.processador: specs.append(f"Proc: {eq.processador}")
        if eq.memoria: specs.append(f"RAM: {eq.memoria}")
        if eq.tamanho_disco: specs.append(f"HD: {eq.tamanho_disco}")
        if eq.sistema_operacional: specs.append(f"OS: {eq.sistema_operacional}")
        if eq.especificacoes: specs.append(eq.especificacoes)
        specs_str = " | ".join(specs) if specs else "N/A"

        row = [
            eq.patrimonio,
            eq.categoria,
            eq.tipo,
            f"{eq.marca} {eq.modelo}",
            eq.numero_serie or "N/A",
            eq.setor or "TI / Estoque",
            eq.responsavel or "Estoque",
            eq.quantidade,
            eq.condicao_inicial or "USADO",
            eq.status.value if hasattr(eq.status, 'value') else str(eq.status),
            valor_val,
            specs_str
        ]
        ws.append(row)
        current_row = ws.max_row
        for col_num in range(1, len(headers) + 1):
            cell = ws.cell(row=current_row, column=col_num)
            cell.font = normal_font
            cell.border = thin_border
            if col_num in [1, 5, 8, 9, 10]:
                cell.alignment = Alignment(horizontal="center")
            if col_num == 11:
                cell.number_format = 'R$ #,##0.00'

    # Summary row
    ws.append([])
    summary_row = [
        "TOTAL DE ATIVOS", "", "", "", "", "", "",
        sum(eq.quantidade for eq in equipamentos), "", "",
        total_valor, ""
    ]
    ws.append(summary_row)
    sum_row_idx = ws.max_row
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=sum_row_idx, column=col_num)
        cell.font = bold_font
        cell.fill = gold_fill
        cell.border = thin_border
        if col_num == 11:
            cell.number_format = 'R$ #,##0.00'

    # Auto-adjust column widths
    for col in ws.columns:
        max_len = 0
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        for cell in col:
            if cell.row > 4 and cell.value:
                max_len = max(max_len, len(str(cell.value)))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 12)

    ws.column_dimensions['L'].width = 40 # Specifications column

    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer

def gerar_pdf_inventario(equipamentos) -> io.BytesIO:
    buffer = io.BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=landscape(letter), # Horizontal page for rich inventory table
        rightMargin=25,
        leftMargin=25,
        topMargin=25,
        bottomMargin=25
    )

    styles = getSampleStyleSheet()
    
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontSize=16,
        leading=20,
        alignment=0,
        textColor=colors.HexColor('#162052'),
        fontName='Helvetica-Bold'
    )

    sub_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontSize=9,
        leading=12,
        alignment=0,
        textColor=colors.HexColor('#EAB308'),
        fontName='Helvetica-Bold'
    )

    table_header_style = ParagraphStyle(
        'TableHeader',
        parent=styles['Normal'],
        fontSize=8,
        leading=10,
        alignment=1,
        textColor=colors.white,
        fontName='Helvetica-Bold'
    )

    cell_style = ParagraphStyle(
        'CellNormal',
        parent=styles['Normal'],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.black,
        fontName='Helvetica'
    )

    cell_bold = ParagraphStyle(
        'CellBold',
        parent=styles['Normal'],
        fontSize=7.5,
        leading=9.5,
        textColor=colors.HexColor('#162052'),
        fontName='Helvetica-Bold'
    )

    story = []

    # Logo & Header
    logo_path = os.path.join(os.path.dirname(__file__), 'logo-dia.png')
    if os.path.exists(logo_path):
        img_logo = Image(logo_path, width=140, height=48)
        header_table = Table([
            [
                img_logo,
                [
                    Paragraph("<b>DIA DISTRIBUIÇÃO E IMPORTAÇÃO AFOGADOS LTDA</b>", title_style),
                    Paragraph("RELATÓRIO EXECUTIVO DE INVENTÁRIO DE TI & ATIVOS", sub_style),
                    Paragraph(f"Emitido em: {datetime.now().strftime('%d/%m/%Y às %H:%M')} | Total de Equipamentos: {len(equipamentos)}", ParagraphStyle('DateSub', parent=cell_style, textColor=colors.HexColor('#64748B')))
                ]
            ]
        ], colWidths=[150, 590])
        header_table.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ]))
        story.append(header_table)
    else:
        story.append(Paragraph("DIA DISTRIBUIÇÃO E IMPORTAÇÃO AFOGADOS LTDA", title_style))

    story.append(Spacer(1, 10))

    # Data Table
    # Width total: 742pt
    table_data = [
        [
            Paragraph("<b>Patrimônio</b>", table_header_style),
            Paragraph("<b>Categoria</b>", table_header_style),
            Paragraph("<b>Tipo / Equipamento</b>", table_header_style),
            Paragraph("<b>Marca / Modelo</b>", table_header_style),
            Paragraph("<b>Nº de Série</b>", table_header_style),
            Paragraph("<b>Setor em Uso</b>", table_header_style),
            Paragraph("<b>Responsável</b>", table_header_style),
            Paragraph("<b>Qtd</b>", table_header_style),
            Paragraph("<b>Status</b>", table_header_style),
            Paragraph("<b>Valor (R$)</b>", table_header_style)
        ]
    ]

    total_qtd = 0
    total_val = 0.0

    for eq in equipamentos:
        total_qtd += eq.quantidade
        v = (eq.valor or 0.0) * eq.quantidade
        total_val += v

        st_str = eq.status.value if hasattr(eq.status, 'value') else str(eq.status)
        
        table_data.append([
            Paragraph(eq.patrimonio, cell_bold),
            Paragraph(eq.categoria, cell_style),
            Paragraph(eq.tipo, cell_style),
            Paragraph(f"{eq.marca} {eq.modelo}", cell_style),
            Paragraph(eq.numero_serie or "N/A", cell_style),
            Paragraph(eq.setor or "TI / Estoque", cell_style),
            Paragraph(eq.responsavel or "Estoque", cell_style),
            Paragraph(str(eq.quantidade), cell_style),
            Paragraph(st_str, cell_bold),
            Paragraph(f"R$ {v:,.2f}".replace(',', 'X').replace('.', ',').replace('X', '.'), cell_style)
        ])

    # Summary Row
    table_data.append([
        Paragraph("<b>TOTAL GERAL</b>", table_header_style),
        "", "", "", "", "", "",
        Paragraph(f"<b>{total_qtd}</b>", table_header_style),
        "",
        Paragraph(f"<b>R$ {total_val:,.2f}</b>".replace(',', 'X').replace('.', ',').replace('X', '.'), table_header_style)
    ])

    col_widths = [75, 75, 85, 95, 85, 90, 95, 30, 55, 57]
    t_inv = Table(table_data, colWidths=col_widths, repeatRows=1)
    t_inv.setStyle(TableStyle([
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#162052')),
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('VALIGN', (0,0), (-1,-1), 'MIDDLE'),
        ('PADDING', (0,0), (-1,-1), 4),
        ('BACKGROUND', (0,-1), (-1,-1), colors.HexColor('#F59E0B')),
        ('ALIGN', (7,0), (7,-1), 'CENTER'),
        ('ALIGN', (9,0), (9,-1), 'RIGHT'),
    ]))
    story.append(t_inv)

    doc.build(story)
    buffer.seek(0)
    return buffer
